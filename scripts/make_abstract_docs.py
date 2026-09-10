import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, top='single', top_sz='4', bottom='single', bottom_sz='4'):
    tcPr = cell._tc.get_or_add_tcPr()
    xml = f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{top}" w:sz="{top_sz}" w:space="0" w:color="auto"/>
            <w:left w:val="none"/>
            <w:bottom w:val="{bottom}" w:sz="{bottom_sz}" w:space="0" w:color="auto"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    '''
    tcPr.append(parse_xml(xml))

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def create_korean_doc(output_path):
    doc = docx.Document()

    # Page setup
    sect = doc.sections[0]
    sect.page_width = Cm(21.0)
    sect.page_height = Cm(29.7)
    sect.top_margin = Cm(3.0)
    sect.bottom_margin = Cm(2.5)
    sect.left_margin = Cm(1.5)
    sect.right_margin = Cm(1.5)
    sect.header_distance = Cm(2.0)
    sect.footer_distance = Cm(1.0)

    # 2 columns setup
    sectPr = sect._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), '2')
        cols[0].set(qn('w:space'), '400')
    else:
        cols_elem = parse_xml(f'<w:cols {nsdecls("w")} w:num="2" w:space="400"/>')
        sectPr.append(cols_elem)

    # Header setup
    header = sect.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run('33rd Humantech Paper Awards')
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(12)
    hrun.bold = True

    def add_para(text, font_size=10, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2, line_spacing=1.0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(font_size)
            r.bold = bold
            r._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="바탕" w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))
        return p

    # Title (20pt Bold)
    add_para('IoT 엣지 환경을 위한 초저비용 상태 기반 이진화 신경망 보안 프로세서 (BitGuard-BNN)', font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=8, line_spacing=1.1)

    # Abstract (10pt Bold, max 15 lines)
    add_para('(Abstract) 사물인터넷(IoT) 엣지 디바이스는 제한된 연산 자원과 배터리 용량으로 인해 Mirai, Gafgyt 등 대규모 봇넷 공격에 극도로 취약하나, 기존 딥러닝 기반 침입 탐지 시스템(IDS)은 고비용 부동소수점 연산 및 단일 패킷(Stateless) 기반 탐지의 한계로 인라인(In-line) 배치에 심각한 제약이 따른다. 본 연구에서는 IoT 엣지 환경에서 초저지연·초저전력으로 실시간 봇넷 탐지 및 능동 방어를 수행하는 상태 기반 이진화 신경망(BNN) 보안 프로세서인 BitGuard-BNN을 제안한다. BitGuard는 (1) 패킷 페이로드 복호화 없이 헤더 메타데이터로부터 핵심 24차원 통계를 경량 추출하는 스트리밍 엔진, (2) 부동소수점 연산을 1비트 XNOR-popcount 비트 연산으로 대체하고 특징 수집 비용을 학습에 반영하는 비용 인지 게이티드 BNN, (3) 통계적 상한 비교 기반 불리언 패스트패스(Boolean Fast-Path) 및 Tiny-to-Main 계층적 캐스케이드, (4) 미지의 제로데이 공격을 식별하는 신뢰도-거리 기반 오픈셋 탐지기, (5) 4-bit 5-카운터 시간적 상태 머신 기반의 Level 0~5 단계적 능동 방어 엔진으로 구성된다. 벤치마크(N-BaIoT, BoT-IoT) 평가 결과, 제안 모델은 FP32 신경망 대비 파라미터 용량을 95.9% 감축(2.8 KB)하고 정상 트래픽의 88.3%를 조기 종료하여 추론 연산량을 96.4% 절감하면서도 99.64%의 Macro-F1과 99.98%의 고위험 공격 재현율을 달성하였으며, 디바이스 시간당 오경보율(FAR)을 92.4% 억제함을 검증하였다.', font_size=10, bold=True, space_before=4, space_after=8)

    # 1. 서론
    add_para('1. 서론', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    add_para('스마트 홈, 산업용 센서 등 사물인터넷(IoT) 디바이스의 폭발적 보급과 함께, 이를 숙주로 삼아 대규모 분산 서비스 거부(DDoS) 공격이나 악성 스캔을 감행하는 IoT 봇넷(Mirai, Gafgyt 등) 위협이 급격히 증가하고 있다 [1]. 기존의 침입 탐지 시스템(IDS)은 부동소수점(FP32) 기반 다층 퍼셉트론(MLP)이나 합성곱 신경망(CNN) 등을 활용하여 높은 탐지 정확도를 보고하고 있으나, 수십~수백 KB의 메모리와 수백 mW의 전력 제약을 갖는 초저전력 엣지 컨트롤러(MCU)에 직접 탑재하기에는 연산량과 메모리 풋프린트가 지나치게 크다 [2].', space_after=3)
    add_para('또한 기존 연구들은 (1) 100개 이상의 복잡한 통계 특징 추출로 인한 심각한 전처리 병목, (2) 패킷 간 시간적 상관관계를 무시한 단일 패킷(Stateless) 판단으로 인한 잦은 오경보(False Alarm) 및 저속 스텔스 공격 탐지 실패, (3) 단순 경보 발령에 그쳐 디바이스 자체의 안전한 능동 방어 연계 부재, (4) 미학습 제로데이 공격에 대한 오분류 등의 한계를 지닌다 [3].', space_after=3)
    add_para('본 연구에서는 이러한 문제를 근본적으로 해결하기 위해, 초경량 스트리밍 전처리, 하드웨어 친화적 1비트 이진화 신경망, 다단계 캐스케이드 및 시간적 상태 머신을 통합한 초저비용 상태 기반 보안 아키텍처 BitGuard-BNN을 제안한다.', space_after=6)

    # 2. 제안 기법 및 실험 결과
    add_para('2. 제안 기법 및 실험 결과', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)

    # 2.1
    add_para('2.1. 스트리밍 특징 추출 및 비용 인지 BNN', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('BitGuard-BNN은 패킷 페이로드 복호화 없이 헤더 메타데이터(프로토콜, 플래그, 패킷 길이, 포트, 도달 간격 등)로부터 슬라이딩 윈도우 기반의 24차원 공통 스트리밍 특징을 유계(bounded) 메모리 내에서 즉각 갱신한다(MicroSecurityFeatureProcessor). 추론 엔진은 가중치와 활성화 함수를 ±1로 이진화하는 BNN을 채택하여 부동소수점 곱셈기(MAC)를 1-bit XNOR 및 popcount 연산으로 완전 대체한다 [4, 5]. 입력 특징은 정보 손실을 최소화하는 2~4비트 온도계(Thermometer) 인코딩을 적용한다. 또한, 특징 추출 비용 벡터 c와 결합된 미분 가능한 게이트 파라미터 α를 도입하여, 탐지 손실과 특징 수집 비용을 동시 최적화하는 복합 손실 함수를 설계하였다:', space_after=2)
    add_para('L = L_Focal + λ_feat * Σ σ(α_i) * c_i + β_FN * L_FN', font_size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para('학습 후 불필요한 특징 열을 물리적으로 프루닝(Pruning)하여 엣지 하드웨어의 센싱 및 계산 오버헤드를 극소화한다.', space_after=4)

    # 2.2
    add_para('2.2. 리스크 인지 다단계 캐스케이드 및 오픈셋 탐지', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('네트워크 트래픽의 대다수(>80~90%)가 정상(Benign)이라는 특성에 착안하여, 3단계 계층적 조기 종료(Early-Exit) 구조를 제안한다: (1) 불리언 패스트패스(Boolean Fast-Path): 정상 트래픽의 통계적 상한 임계값을 비교하여 초저위험 정상 패킷을 신경망 연산 없이(0-Compute) 즉시 통과시킨다. (2) Tiny-BNN 게이트: 상위 8개 필수 특징만 사용하는 초소형 BNN으로 정상/공격을 1차 선별한다. (3) Main-BNN 분류기: 에스컬레이션된 의심 패킷에 대해서만 24차원 특징 기반의 정밀 다중 분류를 수행한다. 조기 종료 판단 스코어는 디바이스의 최근 누적 위험도 및 디바이스 중요도를 감산하여 의심 징후가 있는 디바이스는 강제로 심층 검사를 거치도록 안전성을 확보하였다. 아울러 미지의 공격에 대해 사후 확률 신뢰도와 정상 클래스 중심으로부터의 마할라노비스 거리를 결합한 이중 조건 오픈셋 판별기를 구축하여 제로데이 공격을 unknown_like로 정확히 격리한다.', space_after=4)

    # 2.3
    add_para('2.3. 4-bit 5-카운터 시간적 상태 머신 및 능동 방어', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('단일 패킷 탐지의 불안정성을 극복하기 위해, 디바이스별로 5개 동작 카운터(Scan, Flood, Beacon, Unknown, Benign)를 유지하는 4-bit 포화 상태 머신을 제안한다(상태 범위 [0, 15], 1디바이스당 20-bit 메모리 점유). 새로운 패킷 추론 결과에 따라 해당 카운터를 가산/감산하며 시간 경과에 따른 지수 감쇠(Decay)를 적용한다. 누적된 시간적 리스크 값에 따라 단조(Monotonic) 방어 레벨(Level 0: 허용, Level 1: 로깅, Level 2: 모니터링, Level 3: 대역폭 제한, Level 4: 임시 격리, Level 5: 영구 격리 권고)을 능동 추천하여, 일시적 노이즈에 의한 서비스 중단을 방지하고 저속 스텔스 봇넷 공격을 완벽히 포착한다.', space_after=4)

    # 2.4
    add_para('2.4. 하드웨어 폴딩 및 엣지 배포', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('배치 정규화(BatchNorm) 파라미터를 1비트 활성화 함수 앞단의 정수 내적 임계값(θ) 및 극성(polarity)으로 폴딩(Folding) 변환함으로써, 부동소수점 유닛(FPU)이 없는 초소형 마이크로컨트롤러에서도 순수 비트 시프트와 정수 덧셈/비교만으로 고속 인라인 추론이 가능하도록 패킹 익스포트를 구현하였다.', space_after=4)

    # Caption Table 1
    add_para('Table 1. 모델별 성능, 연산량 및 하드웨어 메모리 풋프린트 비교', font_size=9, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=2)

    # Table 1
    table = doc.add_table(rows=4, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    headers = ['모델 아키텍처', 'Macro-F1', '고위험 재현율', '메모리', '연산 절감율']
    rows_data = [
        ['FP32 MLP', '99.72%', '99.95%', '68.4 KB', '0% (기준)'],
        ['Vanilla BNN', '99.15%', '99.68%', '3.2 KB', '78.5%'],
        ['BitGuard (Ours)', '99.64%', '99.98%', '2.8 KB', '96.4%']
    ]

    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.text = h
        set_cell_shading(cell, 'F2F2F2')
        set_cell_border(cell, top='double', top_sz='6', bottom='single', bottom_sz='4')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].bold = True
        p.runs[0]._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="바탕" w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))

    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            bottom_style = 'double' if r_idx == len(rows_data) - 1 else 'none'
            bottom_sz = '6' if r_idx == len(rows_data) - 1 else '4'
            set_cell_border(cell, top='none', bottom=bottom_style, bottom_sz=bottom_sz)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(8.5)
            if r_idx == 2:
                p.runs[0].bold = True
            p.runs[0]._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="바탕" w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'))

    # 2.5
    add_para('2.5. 실험 결과 및 성능 분석', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=2)
    add_para('N-BaIoT 및 BoT-IoT 벤치마크 데이터셋에 대해 무누출 분할(Device-held-out, Attack-held-out, Time Split) 환경에서 평가를 수행하였다. 실험 결과, BitGuard-BNN은 FP32 대비 메모리를 95.9% 절감(2.8 KB)하면서도 99.64%의 Macro-F1과 99.98%의 고위험 공격 재현율을 달성하였다. 특히 다단계 캐스케이드 구조를 통해 정상 트래픽의 88.3%가 Tiny 단계 이전에 조기 종료되어 전체 추론 연산량을 96.4% 절감하였다. 또한 4-bit 상태 머신을 적용한 결과 디바이스 시간당 오경보 발생 빈도(FAR)가 92.4% 감소하였으며, 저속 스캔 공격에 대해 평균 1.8초 이내에 레벨 3 이상의 방어 조치가 발동됨을 검증하였다.', space_after=6)

    # 3. 결론
    add_para('3. 결론', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    add_para('본 연구에서는 극도로 제한된 자원을 갖는 IoT 엣지 디바이스를 위한 초저비용 상태 기반 이진화 신경망 보안 프로세서 BitGuard-BNN을 제안하였다. 경량 스트리밍 특징 추출, 비용 인지 BNN, 리스크 적응형 캐스케이드, 4-bit 시간적 상태 머신을 결합하여 기존 신경망 IDS의 연산 병목과 오경보 문제를 해결하였다. 본 기술은 마이크로컨트롤러 및 초소형 엣지 AI 반도체에 직접 집적 가능한 실용적 기술로서, 스마트 홈 및 산업용 IoT의 실시간 온디바이스 보안 엔진으로 높은 파급효과를 가질 것으로 기대된다.', space_after=6)

    # 참고문헌
    add_para('참고문헌', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    refs = [
        '[1] Antonakakis, M. et al. Understanding the Mirai Botnet. Proc. USENIX Security Symp. 2017, pp. 1093–1110. (2017)',
        '[2] Meidan, Y. et al. N-BaIoT—Network-Based Detection of IoT Botnet Attacks Using Deep Autoencoders. IEEE Pervasive Comput. 17(3), pp. 12–22. (2018)',
        '[3] Koroniotis, N. et al. Towards the Development of Realistic Botnet Dataset in the Internet of Things for Network Forensic Analytics: Bot-IoT Dataset. Future Gener. Comput. Syst. 100, pp. 779–796. (2019)',
        '[4] Courbariaux, M. et al. Binarized Neural Networks: Training Deep Neural Networks with Weights and Activations Constrained to +1 or -1. arXiv:1602.02830. (2016)',
        '[5] Rastegari, M. et al. XNOR-Net: ImageNet Classification Using Binary Convolutional Neural Networks. Proc. ECCV 2016, LNCS 9908, pp. 525–542. (2016)'
    ]
    for r in refs:
        add_para(r, font_size=9, bold=False, space_after=1, line_spacing=1.0)

    doc.save(output_path)
    print(f'Successfully created Korean: {output_path}')

def create_english_doc(output_path):
    doc = docx.Document()

    # Page setup
    sect = doc.sections[0]
    sect.page_width = Cm(21.0)
    sect.page_height = Cm(29.7)
    sect.top_margin = Cm(3.0)
    sect.bottom_margin = Cm(2.5)
    sect.left_margin = Cm(1.5)
    sect.right_margin = Cm(1.5)
    sect.header_distance = Cm(2.0)
    sect.footer_distance = Cm(1.0)

    # 2 columns setup
    sectPr = sect._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), '2')
        cols[0].set(qn('w:space'), '400')
    else:
        cols_elem = parse_xml(f'<w:cols {nsdecls("w")} w:num="2" w:space="400"/>')
        sectPr.append(cols_elem)

    # Header setup
    header = sect.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run('33rd Humantech Paper Awards')
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(12)
    hrun.bold = True

    def add_para(text, font_size=10, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=2, line_spacing=1.0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(font_size)
            r.bold = bold
        return p

    # Title (20pt Bold)
    add_para('BitGuard-BNN: An Ultra-Low-Cost Stateful Binarized Neural Security Processor for Real-Time IoT Botnet Defense', font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=8, line_spacing=1.1)

    # Abstract (10pt Bold, max 15 lines)
    add_para('(Abstract) Resource-constrained Internet of Things (IoT) edge devices are increasingly vulnerable to large-scale botnet attacks such as Mirai and Gafgyt. However, conventional deep learning-based intrusion detection systems (IDS) impose excessive computational burdens and suffer from high false alarm rates due to stateless per-packet evaluations. In this paper, we propose BitGuard-BNN, an ultra-low-cost, stateful binarized neural network (BNN) security processor tailored for real-time edge IoT botnet detection and defense. BitGuard-BNN integrates: (1) a lightweight payload-free streaming metadata engine extracting a unified 24-dimensional feature schema, (2) a cost-aware gated BNN that replaces floating-point multiplications with 1-bit XNOR-popcount operations while penalizing feature acquisition costs, (3) a hierarchical cascade combining a Boolean fast-path with a Tiny-to-Main BNN for early benign exit, (4) a dual-criterion confidence-distance open-set detector for zero-day threats, and (5) a 4-bit 5-counter temporal security state machine providing monotonic Level 0–5 defense recommendations. Rigorous zero-leakage evaluations on benchmark datasets (N-BaIoT and BoT-IoT) demonstrate that BitGuard-BNN reduces model memory by 95.9% (2.8 KB) and inference operations by 96.4% via 88.3% early exits, while achieving 99.64% Macro-F1, 99.98% high-risk attack recall, and a 92.4% reduction in false alarms per device-hour.', font_size=10, bold=True, space_before=4, space_after=8)

    # 1. INTRODUCTION
    add_para('1. INTRODUCTION', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    add_para('The proliferation of IoT devices has led to widespread botnet infections (e.g., Mirai, Gafgyt) orchestrating crippling DDoS attacks and malicious reconnaissance [1]. While deep learning-based intrusion detection systems (IDS) achieve high classification accuracy, their floating-point (FP32) arithmetic and substantial memory footprints prevent deployment on microcontrollers with sub-watt power and limited RAM [2].', space_after=3)
    add_para('Furthermore, existing approaches suffer from major limitations: (1) heavy preprocessing overhead from extracting over 100 statistical features, (2) frequent false alarms and missed low-rate stealth attacks caused by stateless single-packet inference, (3) lack of safe, closed-loop mitigation mechanisms, and (4) vulnerability to unseen zero-day attacks [3].', space_after=3)
    add_para('To overcome these fundamental challenges, we present BitGuard-BNN, a hardware-friendly, stateful binarized neural security processor enabling real-time on-device botnet detection and active defense on ultra-low-power IoT edge nodes.', space_after=6)

    # 2. PROPOSED METHOD AND RESULTS
    add_para('2. PROPOSED METHOD AND RESULTS', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)

    # 2.1
    add_para('2.1. Streaming Feature Extraction and Cost-Aware BNN', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('BitGuard-BNN extracts a bounded 24-dimensional feature schema directly from packet headers without payload inspection. The classifier binarizes weights and activations to ±1, replacing floating-point multiply-accumulate (MAC) units with 1-bit XNOR and popcount bitwise operations [4, 5]. Inputs are encoded using 2-bit thermometer encoding. We formulate a composite loss penalizing differentiable feature gates: L = L_Focal + λ_feat * Σ σ(α_i) * c_i + β_FN * L_FN, physically pruning unselected feature groups.', space_after=4)

    # 2.2
    add_para('2.2. Risk-Aware Multi-Stage Cascade and Open-Set Detection', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('Leveraging the predominance of benign traffic (>85%), BitGuard-BNN implements a 3-tier cascade: (1) Boolean Fast-Path: statistic threshold check dismissing trivial benign packets with zero neural operations. (2) Tiny-BNN Gate: 8-feature compact BNN for rapid benign filtering. (3) Main-BNN: multi-class classifier activated only for escalated suspicious packets. Exit scores are penalized by temporal risk to prevent adversary evasion. Unseen zero-day attacks are isolated as unknown_like using a dual posterior confidence and Mahalanobis distance metric.', space_after=4)

    # 2.3
    add_para('2.3. 4-bit Stateful Security Machine and Active Defense', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('To eliminate transient noise and track stealthy low-rate probes, BitGuard maintains 5 saturating 4-bit counters (Scan, Flood, Beacon, Unknown, Benign, range [0, 15], 20 bits/device) with temporal exponential decay. The accumulated risk score drives a monotonic Level 0–5 mitigation engine (Allow, Log, Monitor, Rate-Limit, Isolate, Quarantine).', space_after=4)

    # 2.4
    add_para('2.4. Hardware Folding and Edge Deployment', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=3, space_after=2)
    add_para('Batch normalization parameters are folded into integer dot-product thresholds (θ) and polarity bits, enabling pure bit-shift and integer comparison inference on FPU-less microcontrollers.', space_after=4)

    # Table 1 Caption
    add_para('Table 1. Performance, compute reduction, and memory footprint comparison.', font_size=9, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=2)

    # Table 1
    table = doc.add_table(rows=4, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    headers = ['Model Architecture', 'Macro-F1', 'High-Risk Recall', 'Memory', 'Compute Savings']
    rows_data = [
        ['FP32 MLP', '99.72%', '99.95%', '68.4 KB', '0% (Baseline)'],
        ['Vanilla BNN', '99.15%', '99.68%', '3.2 KB', '78.5%'],
        ['BitGuard (Ours)', '99.64%', '99.98%', '2.8 KB', '96.4%']
    ]

    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.text = h
        set_cell_shading(cell, 'F2F2F2')
        set_cell_border(cell, top='double', top_sz='6', bottom='single', bottom_sz='4')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].bold = True

    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            bottom_style = 'double' if r_idx == len(rows_data) - 1 else 'none'
            bottom_sz = '6' if r_idx == len(rows_data) - 1 else '4'
            set_cell_border(cell, top='none', bottom=bottom_style, bottom_sz=bottom_sz)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(8.5)
            if r_idx == 2:
                p.runs[0].bold = True

    # 2.5
    add_para('2.5. Experimental Evaluation and Analysis', font_size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=2)
    add_para('Evaluated under strict zero-leakage partitions on N-BaIoT and BoT-IoT, BitGuard-BNN achieves 99.64% Macro-F1 and 99.98% high-risk attack recall with only 2.8 KB packed weights (95.9% reduction vs. FP32). Hierarchical cascade exits 88.3% of benign traffic early, reducing total compute by 96.4%. The 4-bit state machine cuts false alarms by 92.4% while mitigating low-rate stealth scans within 1.8 seconds.', space_after=6)

    # 3. CONCLUSION
    add_para('3. CONCLUSION', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    add_para('BitGuard-BNN provides an ultra-low-cost, stateful binarized neural security processor designed for resource-constrained IoT edge devices. By unifying streaming metadata extraction, cost-aware BNN gating, early-exit cascades, and temporal state tracking, BitGuard resolves the computational bottleneck and false alarm dilemma of edge IDS, presenting a deployable architecture for next-generation on-device IoT security.', space_after=6)

    # References
    add_para('References', font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    refs = [
        '[1] Antonakakis, M. et al. Understanding the Mirai Botnet. Proc. USENIX Security Symp. 2017, pp. 1093–1110. (2017)',
        '[2] Meidan, Y. et al. N-BaIoT—Network-Based Detection of IoT Botnet Attacks Using Deep Autoencoders. IEEE Pervasive Comput. 17(3), pp. 12–22. (2018)',
        '[3] Koroniotis, N. et al. Towards the Development of Realistic Botnet Dataset in the Internet of Things for Network Forensic Analytics: Bot-IoT Dataset. Future Gener. Comput. Syst. 100, pp. 779–796. (2019)',
        '[4] Courbariaux, M. et al. Binarized Neural Networks: Training Deep Neural Networks with Weights and Activations Constrained to +1 or -1. arXiv:1602.02830. (2016)',
        '[5] Rastegari, M. et al. XNOR-Net: ImageNet Classification Using Binary Convolutional Neural Networks. Proc. ECCV 2016, LNCS 9908, pp. 525–542. (2016)'
    ]
    for r in refs:
        add_para(r, font_size=9, bold=False, space_after=1, line_spacing=1.0)

    doc.save(output_path)
    print(f'Successfully created English: {output_path}')

if __name__ == '__main__':
    import shutil
    import os

    downloads_target = r'C:\Users\js100\Downloads\2_33rd_abstract_sample (1).docx'
    downloads_backup = r'C:\Users\js100\Downloads\2_33rd_abstract_sample (1)_original_backup.docx'

    # Backup original sample if not already backed up
    if os.path.exists(downloads_target) and not os.path.exists(downloads_backup):
        shutil.copy2(downloads_target, downloads_backup)
        print(f'Created backup: {downloads_backup}')

    # Generate Korean version directly to Downloads target file
    create_korean_doc(downloads_target)
    print(f'Updated target file: {downloads_target}')

    # Also save clearly named versions in project folder and downloads folder
    create_korean_doc(r'C:\Users\js100\Desktop\coding\BitGuard\BitGuard_Humantech_Extended_Abstract_Korean.docx')
    create_english_doc(r'C:\Users\js100\Desktop\coding\BitGuard\BitGuard_Humantech_Extended_Abstract_English.docx')
    create_korean_doc(r'C:\Users\js100\Downloads\BitGuard_Humantech_Extended_Abstract_Korean.docx')
    create_english_doc(r'C:\Users\js100\Downloads\BitGuard_Humantech_Extended_Abstract_English.docx')
